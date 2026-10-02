"""Leitura e gravacao da aba "Adquirencia" (preserva as outras abas)."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook

from .metricas import Dia, LinhaCalculada, calcular

ABA = "Adquirencia"
CABECALHO = [
    "Data",
    "Faturamento",
    "Qtd Transações",
    "Ticket Médio",
    "Semana",
    "Faturamento Semana (acum.)",
    "Maior Ticket Médio (até a data)",
    "Maior Faturamento Semanal (até a data)",
    "Atualizado em",
]
FORMATOS = {2: "#,##0.00", 3: "#,##0", 4: "#,##0.00", 6: "#,##0.00", 7: "#,##0.00", 8: "#,##0.00", 1: "dd/mm/yyyy", 9: "dd/mm/yyyy hh:mm"}


def abrir(conteudo: bytes | None) -> Workbook:
    return load_workbook(BytesIO(conteudo)) if conteudo else Workbook()


def _aba(wb: Workbook):
    if ABA in wb.sheetnames:
        return wb[ABA]
    # um workbook novo traz uma aba "Sheet" vazia: reaproveita em vez de deixar sobrando
    if wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        ws = wb["Sheet"]
        ws.title = ABA
        return ws
    return wb.create_sheet(ABA)


def ler_historico(wb: Workbook) -> dict[date, Dia]:
    if ABA not in wb.sheetnames:
        return {}
    historico: dict[date, Dia] = {}
    for data, fat, qtd, *_ in wb[ABA].iter_rows(min_row=2, max_col=3, values_only=True):
        if data is None or fat is None:
            continue
        d = data.date() if isinstance(data, datetime) else data
        historico[d] = Dia(d, Decimal(str(fat)), int(qtd or 0))
    return historico


def aplicar(wb: Workbook, novos: list[Dia], agora: datetime) -> list[LinhaCalculada]:
    """Mescla os dias consultados no historico (por data) e regrava a aba.

    Rodar duas vezes no mesmo dia nao duplica linha: a data e a chave.
    """
    historico = ler_historico(wb)
    for dia in novos:
        historico[dia.data] = dia
    linhas = calcular(list(historico.values()))

    ws = _aba(wb)
    ws.delete_rows(1, ws.max_row)
    ws.append(CABECALHO)
    for l in linhas:
        ws.append(
            [
                l.data,
                float(l.faturamento),
                l.qtd_transacoes,
                float(l.ticket_medio),
                l.semana,
                float(l.faturamento_semana),
                float(l.maior_ticket_medio),
                float(l.maior_faturamento_semana),
                agora,
            ]
        )
    for coluna, formato in FORMATOS.items():
        for celula in ws.iter_rows(min_row=2, min_col=coluna, max_col=coluna):
            celula[0].number_format = formato
    ws.freeze_panes = "A2"
    for i, titulo in enumerate(CABECALHO, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(titulo) + 2)
    return linhas


def salvar(wb: Workbook) -> bytes:
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
